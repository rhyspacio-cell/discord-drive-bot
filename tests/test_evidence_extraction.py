from io import BytesIO

from docx import Document

from modules.drive import _collect_documents_for_query
from modules.evidence import ExtractionStatus
from modules.llm import answer_drive_question
from modules.query_constraints import extract_query_constraints


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class FakeDriveService:
    def __init__(self, files):
        self.files_data = files

    def files(self):
        return self

    def list(self, **_kwargs):
        return self

    def execute(self):
        return {"files": self.files_data}


def make_docx_bytes():
    document = Document()
    document.add_paragraph("ACTIVITY RECORD")
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Person"
    table.cell(0, 1).text = "John Smith"
    table.cell(1, 0).text = "Activity"
    table.cell(1, 1).text = "Equipment Inspection"
    table.cell(2, 0).text = "Activity Date"
    table.cell(2, 1).text = "October 10, 2026"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def one_candidate(file_id="drive-doc-1", name="activity_record.docx"):
    return {
        "id": file_id,
        "name": name,
        "mimeType": DOCX_MIME,
        "modifiedTime": "2026-09-29T00:00:00Z",
        "size": "1000",
    }


def query_plan(question):
    return {"query_constraints": extract_query_constraints(question)}


def collect(monkeypatch, file_infos, extractor):
    if extractor is not None:
        monkeypatch.setattr("modules.drive.extract_file_text", extractor)
    return _collect_documents_for_query(
        FakeDriveService(file_infos),
        "test query",
        "What is the planned activity date for John Smith?",
        query_plan("What is the planned activity date for John Smith?"),
    )


def test_docx_table_text_reaches_evidence_assessment(monkeypatch):
    raw_docx = make_docx_bytes()
    monkeypatch.setattr("modules.drive.download_drive_file", lambda *_args: raw_docx)
    file_info = one_candidate()

    documents, _skipped, audit = collect(monkeypatch, [file_info], None)
    text = documents[0]["text"]

    assert "John Smith" in text
    assert "Equipment Inspection" in text
    assert "October 10, 2026" in text
    assert documents[0]["text"] == text
    assessment = audit["assessment_records"][0]
    assert assessment["file_id"] == "drive-doc-1"
    assert assessment["candidate_id"] == "drive-doc-1"
    assert assessment["extraction_status"] == ExtractionStatus.SUCCESS.value
    assert assessment["extracted_char_count"] == len(text)
    assert assessment["extracted_text_available"] is True
    assert assessment["subject_match"] is True
    assert assessment["requested_field_match"] is True
    assert assessment["eligible"] is True


def test_empty_extraction_is_distinct_and_never_reaches_answer_llm(monkeypatch):
    documents, _skipped, audit = collect(
        monkeypatch,
        [one_candidate()],
        lambda _service, file_info: (file_info["name"], ""),
    )
    monkeypatch.setattr(
        "modules.llm.generate_local_response",
        lambda _prompt: (_ for _ in ()).throw(AssertionError("LLM must not run")),
    )

    answer = answer_drive_question(
        "What is the planned activity date for John Smith?",
        documents,
        query_plan("What is the planned activity date for John Smith?"),
        audit,
    )

    assert documents == []
    assert audit["assessment_records"][0]["extraction_status"] == ExtractionStatus.EMPTY.value
    assert audit["assessment_records"][0]["extracted_char_count"] == 0
    assert "couldn't read" in answer.lower()


def test_extraction_exception_is_distinct_and_keeps_error(monkeypatch):
    def fail_extract(_service, _file_info):
        raise RuntimeError("synthetic DOCX parse failure")

    documents, _skipped, audit = collect(monkeypatch, [one_candidate()], fail_extract)

    assert documents == []
    diagnostic = audit["candidate_diagnostics"][0]
    assessment = audit["assessment_records"][0]
    assert diagnostic["extraction_status"] == ExtractionStatus.FAILURE.value
    assert "synthetic DOCX parse failure" in diagnostic["extraction_error"]
    assert assessment["extraction_status"] == ExtractionStatus.FAILURE.value
    assert assessment["extraction_error"] == diagnostic["extraction_error"]

    monkeypatch.setattr(
        "modules.llm.generate_local_response",
        lambda _prompt: (_ for _ in ()).throw(AssertionError("LLM must not run")),
    )
    answer = answer_drive_question(
        "What is the planned activity date for John Smith?",
        documents,
        query_plan("What is the planned activity date for John Smith?"),
        audit,
    )
    assert "couldn't read" in answer.lower()


def test_unrelated_extraction_failure_does_not_hide_readable_candidate(monkeypatch):
    monkeypatch.setattr(
        "modules.llm.generate_local_response",
        lambda _prompt: (_ for _ in ()).throw(AssertionError("No evidence should reach the LLM")),
    )
    audit = {
        "candidates": ["activity_record.docx", "oversized.pdf"],
        "analyzed": ["activity_record.docx"],
        "failed": [{"name": "oversized.pdf", "reason": "download limit"}],
        "candidate_diagnostics": [
            {
                "file_name": "activity_record.docx",
                "candidate_id": "drive-doc-1",
                "extraction_status": ExtractionStatus.SUCCESS.value,
                "extracted_char_count": 281,
                "extracted_text_available": True,
            }
        ],
    }

    answer = answer_drive_question(
        "What date is the equipment inspection planned for John Smith?",
        [],
        search_audit=audit,
    )

    assert "readable files" in answer.lower()
    assert "couldn't read" not in answer.lower()
    assert audit["state"] == "candidates_found_but_no_valid_evidence"


def test_short_successful_document_is_not_rejected_for_length(monkeypatch):
    short_text = "Person: John Smith. Activity Date: October 10, 2026."
    documents, _skipped, audit = collect(
        monkeypatch,
        [one_candidate()],
        lambda _service, file_info: (file_info["name"], short_text),
    )

    assert documents[0]["text"] == short_text
    assessment = audit["assessment_records"][0]
    assert assessment["extraction_status"] == ExtractionStatus.SUCCESS.value
    assert assessment["extracted_char_count"] == len(short_text)
    assert assessment["eligible"] is True


def test_representative_is_not_bound_as_activity_subject():
    question = "What is the planned activity date for Maria Santos?"
    text = (
        "Person | John Smith\n"
        "Representative | Maria Santos\n"
        "Activity | Equipment Inspection\n"
        "Activity Date | October 10, 2026"
    )
    from modules.evidence import assess_candidate

    assessment = assess_candidate(
        question,
        {"name": "activity_record.docx", "text": text},
        {"query_constraints": extract_query_constraints(question)},
    )

    assert assessment.extraction_status is ExtractionStatus.SUCCESS
    assert assessment.subject_match is False
    assert assessment.relationship_supported is False
    assert assessment.eligible is False


def test_entity_free_activity_query_accepts_structured_activity_and_date():
    question = "When is the equipment inspection scheduled?"
    text = (
        "Activity | Equipment Inspection\n"
        "Activity Date | October 10, 2026"
    )
    from modules.evidence import assess_candidate

    assessment = assess_candidate(
        question,
        {"name": "activity_record.docx", "text": text},
        {"query_constraints": extract_query_constraints(question)},
    )

    assert assessment.subject_match is None
    assert assessment.activity_match is True
    assert assessment.requested_field_match is True
    assert assessment.relationship_supported is True
    assert assessment.eligible is True


def test_duplicate_filenames_keep_separate_drive_id_assessments(monkeypatch):
    question = "Summarize the project risk assessment."
    files = [
        one_candidate("drive-id-a", "project_risk_assessment.docx"),
        one_candidate("drive-id-b", "project_risk_assessment.docx"),
    ]
    texts = {
        "drive-id-a": "Project risk assessment identifies schedule risk.",
        "drive-id-b": "Project risk assessment identifies cost risk.",
    }
    monkeypatch.setattr(
        "modules.drive.extract_file_text",
        lambda _service, file_info: (file_info["name"], texts[file_info["id"]]),
    )
    documents, _skipped, audit = _collect_documents_for_query(
        FakeDriveService(files),
        "test query",
        question,
        {"query_constraints": {}},
    )

    assert len(documents) == 2
    assert {doc["file_id"] for doc in documents} == {"drive-id-a", "drive-id-b"}
    assert {record["file_id"] for record in audit["assessment_records"]} == {
        "drive-id-a",
        "drive-id-b",
    }
    assert {doc["text"] for doc in documents} == set(texts.values())


def test_unsupported_candidate_has_status_and_never_becomes_evidence():
    file_info = {
        "id": "unsupported-id",
        "name": "archive.bin",
        "mimeType": "application/octet-stream",
        "modifiedTime": "2026-09-29T00:00:00Z",
    }
    documents, _skipped, audit = _collect_documents_for_query(
        FakeDriveService([file_info]),
        "test query",
        "Summarize the project risk assessment.",
        {"query_constraints": {}},
    )

    assert documents == []
    diagnostic = audit["candidate_diagnostics"][0]
    assessment = audit["assessment_records"][0]
    assert diagnostic["extraction_status"] == ExtractionStatus.UNSUPPORTED.value
    assert assessment["extraction_status"] == ExtractionStatus.UNSUPPORTED.value
    assert assessment["eligible"] is False


def test_equal_score_candidates_use_stable_name_and_id_order(monkeypatch):
    files = [
        {
            "id": "drive-b",
            "name": "activity_record_b.txt",
            "mimeType": "text/plain",
            "modifiedTime": "2026-09-29T00:00:00Z",
        },
        {
            "id": "drive-a",
            "name": "activity_record_a.txt",
            "mimeType": "text/plain",
            "modifiedTime": "2026-09-29T00:00:00Z",
        },
    ]
    monkeypatch.setattr(
        "modules.drive.score_file_relevance",
        lambda *_args, **_kwargs: 1,
    )
    monkeypatch.setattr(
        "modules.drive.extract_file_text",
        lambda _service, file_info: (
            file_info["name"],
            "Person: John Smith\nActivity Date: October 10, 2026",
        ),
    )

    documents, _skipped, _audit = _collect_documents_for_query(
        FakeDriveService(files),
        "query",
        "What is the planned activity date for John Smith?",
        {"query_constraints": extract_query_constraints(
            "What is the planned activity date for John Smith?"
        )},
    )

    assert [document["name"] for document in documents] == [
        "activity_record_a.txt",
        "activity_record_b.txt",
    ]
